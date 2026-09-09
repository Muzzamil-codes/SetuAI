import os
import json
from typing import Dict, Any

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


def generate_pptx(data: dict, output_dir: str = "outputs") -> Dict[str, Any]:
    """Generate a .pptx presentation from structured data.
    
    Args:
        data: dict with keys: title, company, reviewer, department, findings, recommendations
        output_dir: directory to write the output file
    
    Returns:
        dict with keys: success, filename, path, error
    """
    os.makedirs(output_dir, exist_ok=True)
    
    if not HAS_PPTX:
        return {"success": False, "filename": "", "path": "", "error": "python-pptx not installed"}
    
    try:
        prs = Presentation()
        
        # Slide 1: Title
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = data.get("title", "Presentation")
        slide.placeholders[1].text = (
            f"Company: {data.get('company', 'N/A')}\n"
            f"Reviewer: {data.get('reviewer', 'N/A')}\n"
            f"Department: {data.get('department', 'N/A')}"
        )
        
        # Slide 2: Executive Summary
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Executive Summary"
        findings = data.get("findings", [])
        slide.placeholders[1].text = f"Total Findings: {len(findings)}"
        
        # Slide 3: Recommendations
        recommendations = data.get("recommendations", [])
        if recommendations:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "Recommendations"
            text_box = slide.placeholders[1]
            text_box.text = ""
            for rec in recommendations:
                text_box.text += f"\u2022 {rec}\n"
        
        # Slide 4: Chart if available
        try:
            from artifact_factory.chart_generator import generate_chart
            chart_path = generate_chart(data, output_dir)
            if chart_path and os.path.exists(chart_path):
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = "Status Chart"
                slide.shapes.add_picture(chart_path, 1000000, 1000000, width=4000000)
        except Exception:
            pass
        
        from datetime import datetime
        filename = f"presentation_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
        path = os.path.join(output_dir, filename)
        prs.save(path)
        
        return {"success": True, "filename": filename, "path": path, "error": None}
    except Exception as e:
        return {"success": False, "filename": "", "path": "", "error": str(e)}


if __name__ == "__main__":
    with open("data.json", "r") as f:
        data = json.load(f)
    result = generate_pptx(data, "output")
    print(f"PPTX Generated: {result}")