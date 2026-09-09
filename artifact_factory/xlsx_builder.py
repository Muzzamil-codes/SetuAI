import os
import json
from typing import Dict, Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def generate_xlsx(data: dict, output_dir: str = "outputs") -> Dict[str, Any]:
    """Generate a professional .xlsx dashboard from structured data.
    
    Args:
        data: dict with keys: title, company, reviewer, department, findings
        output_dir: directory to write the output file
    
    Returns:
        dict with keys: success, filename, path, error
    """
    os.makedirs(output_dir, exist_ok=True)
    
    if not HAS_OPENPYXL:
        # Fallback to CSV
        from datetime import datetime
        filename = f"dashboard_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        path = os.path.join(output_dir, filename)
        with open(path, "w") as f:
            f.write("Field,Value,Status\n")
            for item in data.get("findings", []):
                f.write(f"{item.get('field','')},{item.get('value','')},{item.get('status','')}\n")
        return {"success": True, "filename": filename, "path": path, "error": None}
    
    try:
        wb = Workbook()
        
        # Dashboard Sheet
        dashboard = wb.active
        dashboard.title = "Dashboard"
        dashboard.merge_cells("A1:B1")
        dashboard["A1"] = data.get("title", "Report")
        dashboard["A1"].font = Font(size=20, bold=True)
        dashboard["A1"].alignment = Alignment(horizontal="center")
        
        dashboard["A3"] = "Company"
        dashboard["B3"] = data.get("company", "N/A")
        dashboard["A4"] = "Reviewer"
        dashboard["B4"] = data.get("reviewer", "N/A")
        dashboard["A5"] = "Department"
        dashboard["B5"] = data.get("department", "N/A")
        
        approved = review = rejected = 0
        for item in data.get("findings", []):
            status = item.get("status", "")
            if status == "Approved": approved += 1
            elif status == "Review": review += 1
            elif status == "Rejected": rejected += 1
        
        dashboard["A7"] = "Approved"
        dashboard["B7"] = approved
        dashboard["A8"] = "Review"
        dashboard["B8"] = review
        dashboard["A9"] = "Rejected"
        dashboard["B9"] = rejected
        
        green_fill = PatternFill(fill_type="solid", fgColor="92D050")
        yellow_fill = PatternFill(fill_type="solid", fgColor="FFD966")
        red_fill = PatternFill(fill_type="solid", fgColor="FF6666")
        dashboard["B7"].fill = green_fill
        dashboard["B8"].fill = yellow_fill
        dashboard["B9"].fill = red_fill
        
        # Findings Sheet
        findings_sheet = wb.create_sheet("Findings")
        headers = ["Field", "Value", "Status"]
        findings_sheet.append(headers)
        
        header_fill = PatternFill(fill_type="solid", fgColor="4F81BD")
        for cell in findings_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        
        for item in data.get("findings", []):
            findings_sheet.append([
                item.get("field", ""),
                item.get("value", ""),
                item.get("status", "")
            ])
        
        for column in findings_sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            findings_sheet.column_dimensions[column_letter].width = max_length + 5
        
        from datetime import datetime
        filename = f"dashboard_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        path = os.path.join(output_dir, filename)
        wb.save(path)
        
        return {"success": True, "filename": filename, "path": path, "error": None}
    except Exception as e:
        return {"success": False, "filename": "", "path": "", "error": str(e)}


if __name__ == "__main__":
    with open("data.json", "r") as f:
        data = json.load(f)
    result = generate_xlsx(data, "output")
    print(f"XLSX Generated: {result}")