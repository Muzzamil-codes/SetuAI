import os
import re
import json
from datetime import datetime
from typing import Dict, Any, List, Optional

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def _set_cell_background(cell, fill_hex: str):
    """Set the background color of a table cell."""
    try:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill_hex)
        tcPr.append(shd)
    except Exception:
        pass


def _add_markdown_content(doc: Any, text: str):
    """Parses markdown text into native python-docx headings, lists, bold/italics, and tables."""
    if not text:
        return

    def _add_inline_runs(paragraph, line_text):
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', line_text)
        for part in parts:
            if not part:
                continue
            if part.startswith('**') and part.endswith('**') and len(part) >= 4:
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
                run = paragraph.add_run(part[1:-1])
                try:
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                except Exception:
                    pass
            else:
                paragraph.add_run(part)

    lines = text.split('\n')
    i = 0
    table_buffer = []

    def flush_table(buffer):
        if not buffer or not HAS_DOCX:
            return
        parsed_rows = []
        for raw_line in buffer:
            cols = [c.strip() for c in raw_line.strip().strip('|').split('|')]
            if cols and not all(re.match(r'^:?-+:?$', c) for c in cols):
                parsed_rows.append(cols)
        if not parsed_rows:
            return
        
        max_cols = max(len(r) for r in parsed_rows)
        table = doc.add_table(rows=len(parsed_rows), cols=max_cols)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, row in enumerate(parsed_rows):
            for c_idx in range(max_cols):
                val = row[c_idx] if c_idx < len(row) else ""
                cell = table.cell(r_idx, c_idx)
                cell.text = val
                if r_idx == 0:
                    _set_cell_background(cell, "2B4C7E")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(4)

    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.strip()

        # Handle markdown table lines
        if line.startswith('|') and line.endswith('|'):
            table_buffer.append(line)
            i += 1
            continue
        else:
            if table_buffer:
                flush_table(table_buffer)
                table_buffer = []

        if not line:
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,6})\s+(.*)', line)
        if heading_match:
            level = min(len(heading_match.group(1)), 4)
            h_text = heading_match.group(2)
            p = doc.add_heading(level=level)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            _add_inline_runs(p, h_text)
            i += 1
            continue

        # Bullet lists
        list_match = re.match(r'^([\-\*]|\d+\.)\s+(.*)', line)
        if list_match:
            bullet_type = list_match.group(1)
            b_text = list_match.group(2)
            style = 'List Number' if bullet_type[-1] == '.' else 'List Bullet'
            try:
                p = doc.add_paragraph(style=style)
            except Exception:
                p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _add_inline_runs(p, b_text)
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        _add_inline_runs(p, line)
        i += 1

    if table_buffer:
        flush_table(table_buffer)


def generate_docx(data: dict, output_dir: str = "outputs") -> Dict[str, Any]:
    """Generate a professional Word document (.docx) based on document type.
    
    Supports:
      - 'notice': Official facility/company notice with header, badge, subject, body, directives, sign-off.
      - 'letter'/'memo': Official correspondence with letterhead, recipient, subject, body, closing.
      - 'report'/'approval_note': Technical report or approval note with executive summary, findings, recommendations.
      - 'standard': General structured document.
    """
    os.makedirs(output_dir, exist_ok=True)
    
    title = data.get("title", "Document")
    doc_type = data.get("doc_type", "").lower()
    
    # Auto-detect doc_type if not explicitly set
    title_lower = title.lower()
    if not doc_type:
        if "notice" in title_lower or "announcement" in title_lower or "alert" in title_lower:
            doc_type = "notice"
        elif "letter" in title_lower or "memo" in title_lower or "memorandum" in title_lower:
            doc_type = "letter"
        elif "approval" in title_lower or "report" in title_lower:
            doc_type = "report"
        else:
            doc_type = "standard"

    company = data.get("company_or_org", data.get("company", "Setu Industrial Corporation"))
    department = data.get("department", "Operations & Maintenance Division")
    date_str = data.get("date", datetime.now().strftime('%d %B %Y'))
    recipient = data.get("recipient_or_target", data.get("recipient", "All Concerned Staff & Department Leads"))
    body_content = data.get("body_content", data.get("content", ""))
    signatory = data.get("signatory", data.get("reviewer", "Issuing Authority"))
    action_items = data.get("action_items_or_recommendations", data.get("recommendations", []))
    findings = data.get("findings", [])

    safe_title = re.sub(r'[^a-zA-Z0-9_\-]', '_', title[:30]).strip('_').lower()
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"{doc_type}_{safe_title}_{timestamp}.docx"
    path = os.path.join(output_dir, filename)

    if not HAS_DOCX:
        # Fallback to structured text file
        txt_filename = filename.replace(".docx", ".txt")
        txt_path = os.path.join(output_dir, txt_filename)
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(f"{company.upper()}\n{department}\n{'='*50}\n")
            f.write(f"DOCUMENT TYPE: {doc_type.upper()}\n")
            f.write(f"TITLE: {title}\nDATE: {date_str}\nTO: {recipient}\n\n")
            f.write(f"CONTENT:\n{body_content}\n\n")
            if action_items:
                f.write("ACTION ITEMS / PRECAUTIONS:\n")
                for item in action_items:
                    f.write(f"  - {item}\n")
            f.write(f"\nISSUED BY: {signatory}\n")
        return {"success": True, "filename": txt_filename, "path": txt_path, "doc_type": doc_type, "error": None}

    try:
        doc = Document()

        # Set standard document margins (1 inch)
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Style normal font
        try:
            normal_style = doc.styles['Normal']
            normal_style.font.name = 'Calibri'
            normal_style.font.size = Pt(11)
            normal_style.font.color.rgb = RGBColor(33, 37, 41)
        except Exception:
            pass

        # -------------------------------------------------------------------
        # 1. NOTICE LAYOUT
        # -------------------------------------------------------------------
        if doc_type == "notice":
            # Organization Header
            org_p = doc.add_paragraph()
            org_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            org_p.paragraph_format.space_after = Pt(2)
            run_org = org_p.add_run(company.upper())
            run_org.bold = True
            run_org.font.size = Pt(15)
            run_org.font.color.rgb = RGBColor(27, 54, 93)

            dept_p = doc.add_paragraph()
            dept_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            dept_p.paragraph_format.space_after = Pt(14)
            run_dept = dept_p.add_run(department)
            run_dept.font.size = Pt(10.5)
            run_dept.font.italic = True
            run_dept.font.color.rgb = RGBColor(100, 110, 120)

            # Prominent NOTICE Banner
            badge_p = doc.add_paragraph()
            badge_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            badge_p.paragraph_format.space_after = Pt(16)
            run_badge = badge_p.add_run("OFFICIAL NOTICE")
            run_badge.bold = True
            run_badge.font.size = Pt(16)
            run_badge.font.color.rgb = RGBColor(180, 40, 40)

            # Metadata Table (Ref, Date, Target)
            meta_table = doc.add_table(rows=2, cols=2)
            meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            meta_table.autofit = True
            for r in meta_table.rows:
                for c in r.cells:
                    c.width = Inches(3.2)

            meta_table.cell(0, 0).paragraphs[0].add_run(f"Ref: SETU/NOT/{datetime.now().year}/{datetime.now().strftime('%m%d')}").font.size = Pt(9.5)
            meta_table.cell(0, 1).paragraphs[0].add_run(f"Date: {date_str}").font.size = Pt(9.5)
            meta_table.cell(0, 1).paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT

            meta_table.cell(1, 0).paragraphs[0].add_run(f"Target / Audience: {recipient}").font.size = Pt(9.5)
            meta_table.cell(1, 1).paragraphs[0].add_run("Status: Immediate Action").font.size = Pt(9.5)
            meta_table.cell(1, 1).paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT

            doc.add_paragraph().paragraph_format.space_after = Pt(6)

            # Subject Line
            subj_p = doc.add_paragraph()
            subj_p.paragraph_format.space_before = Pt(8)
            subj_p.paragraph_format.space_after = Pt(10)
            run_subj_label = subj_p.add_run("Subject: ")
            run_subj_label.bold = True
            run_subj_label.font.size = Pt(12)
            run_subj_label.font.color.rgb = RGBColor(27, 54, 93)
            run_subj_val = subj_p.add_run(title)
            run_subj_val.bold = True
            run_subj_val.font.size = Pt(12)

            # Body Content
            _add_markdown_content(doc, body_content)

            # Action Items / Key Directives (if provided)
            if action_items:
                act_h = doc.add_heading(level=2)
                act_h.paragraph_format.space_before = Pt(12)
                act_h.paragraph_format.space_after = Pt(4)
                act_h.add_run("Action Items & Key Directives").font.color.rgb = RGBColor(27, 54, 93)
                for item in action_items:
                    p = doc.add_paragraph(style='List Bullet')
                    p.paragraph_format.space_after = Pt(2)
                    p.add_run(str(item))

            # Sign-off Block
            sig_p = doc.add_paragraph()
            sig_p.paragraph_format.space_before = Pt(24)
            sig_p.paragraph_format.space_after = Pt(2)
            sig_p.add_run("Issued by:").font.size = Pt(10.5)

            sig_name = doc.add_paragraph()
            sig_name.paragraph_format.space_after = Pt(2)
            run_name = sig_name.add_run(signatory)
            run_name.bold = True
            run_name.font.size = Pt(11)

            sig_dept = doc.add_paragraph()
            sig_dept.paragraph_format.space_after = Pt(2)
            sig_dept.add_run(department).font.size = Pt(10)

            sig_org = doc.add_paragraph()
            sig_org.paragraph_format.space_after = Pt(0)
            sig_org.add_run(company).font.size = Pt(10)

        # -------------------------------------------------------------------
        # 2. LETTER / MEMORANDUM LAYOUT
        # -------------------------------------------------------------------
        elif doc_type in ["letter", "memo"]:
            header_p = doc.add_paragraph()
            header_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run_org = header_p.add_run(company.upper())
            run_org.bold = True
            run_org.font.size = Pt(14)
            run_org.font.color.rgb = RGBColor(27, 54, 93)

            dept_p = doc.add_paragraph()
            dept_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            dept_p.paragraph_format.space_after = Pt(14)
            dept_p.add_run(department).font.italic = True

            # Date and Recipient
            meta_p = doc.add_paragraph()
            meta_p.paragraph_format.space_after = Pt(6)
            meta_p.add_run(f"Date: {date_str}\n").font.size = Pt(10.5)
            meta_p.add_run(f"To: {recipient}\n").font.size = Pt(10.5)
            meta_p.add_run(f"From: {signatory}, {department}\n").font.size = Pt(10.5)

            # Subject
            subj_p = doc.add_paragraph()
            subj_p.paragraph_format.space_before = Pt(6)
            subj_p.paragraph_format.space_after = Pt(10)
            subj_p.add_run("Subject: ").bold = True
            subj_p.add_run(title).bold = True

            # Content
            _add_markdown_content(doc, body_content)

            # Action Items if any
            if action_items:
                doc.add_heading("Next Steps & Recommendations", level=2)
                for item in action_items:
                    doc.add_paragraph(str(item), style='List Bullet')

            # Sign-off
            sig_p = doc.add_paragraph()
            sig_p.paragraph_format.space_before = Pt(20)
            sig_p.add_run("Sincerely,\n\n\n").font.size = Pt(11)
            sig_p.add_run(f"{signatory}\n{department}\n{company}").bold = True

        # -------------------------------------------------------------------
        # 3. REPORT / APPROVAL NOTE / STANDARD
        # -------------------------------------------------------------------
        else:
            # Document Title
            title_p = doc.add_paragraph()
            title_p.paragraph_format.space_before = Pt(10)
            title_p.paragraph_format.space_after = Pt(4)
            run_title = title_p.add_run(title)
            run_title.bold = True
            run_title.font.size = Pt(18)
            run_title.font.color.rgb = RGBColor(27, 54, 93)

            # Sub-header meta
            meta_p = doc.add_paragraph()
            meta_p.paragraph_format.space_after = Pt(14)
            meta_p.add_run(f"{company} | {department} | {date_str}").font.color.rgb = RGBColor(110, 120, 130)

            # Body content
            _add_markdown_content(doc, body_content)

            # Findings Table (ONLY if findings actually provided!)
            if findings and len(findings) > 0:
                doc.add_heading("Findings & Inspections", level=2)
                table = doc.add_table(rows=1, cols=3)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                hdr = table.rows[0].cells
                hdr[0].text = "Item / Field"
                hdr[1].text = "Value / Observation"
                hdr[2].text = "Status"
                for c in hdr:
                    _set_cell_background(c, "2B4C7E")
                    for p in c.paragraphs:
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)

                for item in findings:
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(item.get("field", ""))
                    row_cells[1].text = str(item.get("value", ""))
                    row_cells[2].text = str(item.get("status", ""))

            # Recommendations / Action Items
            if action_items and len(action_items) > 0:
                doc.add_heading("Recommendations & Directives", level=2)
                for item in action_items:
                    doc.add_paragraph(str(item), style='List Bullet')

            # Signatory
            sig_p = doc.add_paragraph()
            sig_p.paragraph_format.space_before = Pt(20)
            sig_p.add_run("Approved & Issued by:\n\n").font.size = Pt(10.5)
            sig_p.add_run(f"{signatory}\n{department}\n{company}").bold = True

        doc.save(path)
        return {
            "success": True,
            "filename": filename,
            "path": path,
            "doc_type": doc_type,
            "error": None
        }

    except Exception as e:
        return {"success": False, "filename": "", "path": "", "doc_type": doc_type, "error": str(e)}


if __name__ == "__main__":
    test_data = {
        "doc_type": "notice",
        "title": "Scheduled Replacement of Rusted Pipelines",
        "company_or_org": "Apex Chemicals Plant No. 4",
        "department": "Plant Maintenance Division",
        "date": "10 September 2026",
        "recipient_or_target": "All Facility Supervisors and Shift In-Charges",
        "body_content": (
            "### 1. Purpose & Scope\n"
            "This notice serves to inform all personnel that critical pipeline replacement works "
            "will commence on **Monday, 15 September 2026** at 08:00 hrs.\n\n"
            "### 2. Affected Locations\n"
            "- Main cooling line segment B-12\n"
            "- Secondary effluent conduit C-04\n\n"
            "### 3. Safety Precautions\n"
            "All personnel entering Area 3 must wear mandatory Level-2 PPE."
        ),
        "action_items_or_recommendations": [
            "Isolate line feed valves 2 hours prior to scheduled work.",
            "Verify complete atmospheric gas clearance before hot work."
        ],
        "signatory": "Operations Director"
    }
    res = generate_docx(test_data, "outputs")
    print("Test Generation Result:", res)